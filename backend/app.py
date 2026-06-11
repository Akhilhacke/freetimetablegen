from fastapi import FastAPI, Request, Form, HTTPException, UploadFile, File
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from database import (init_db, migrate_db, get_db, hash_password, verify_password, 
                      create_session, get_user_from_token)
from timetable_engine import TimetableGenerator, DAYS
import os
import io
import json
import csv
import secrets

app = FastAPI(title="FreeTimetableGen - AI Timetable Generator")

BASE_DIR = os.path.dirname(os.path.dirname(__file__))
app.mount("/static", StaticFiles(directory=os.path.join(BASE_DIR, "static")), name="static")
templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "templates"))

@app.on_event("startup")
def startup():
    init_db()
    migrate_db()

def get_current_user(request: Request):
    token = request.cookies.get("session_token")
    return get_user_from_token(token)

def require_admin(request: Request):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(status_code=403, detail="Admin access required")
    return user

# ─── AUTH ROUTES ───
@app.get("/", response_class=HTMLResponse)
async def landing(request: Request):
    return templates.TemplateResponse(request, "landing.html")

@app.get("/login", response_class=HTMLResponse)
async def login_page(request: Request):
    return templates.TemplateResponse(request, "login.html")

@app.post("/login")
async def login(request: Request, username: str = Form(...), password: str = Form(...)):
    db = get_db()
    user = db.execute("SELECT * FROM users WHERE username=? OR email=?", (username, username)).fetchone()
    db.close()
    if not user or not verify_password(password, user['password_hash']):
        return templates.TemplateResponse(request, "login.html", {"error": "Invalid credentials"})
    token = create_session(user['id'])
    response = RedirectResponse(url="/dashboard", status_code=303)
    response.set_cookie("session_token", token, httponly=True, max_age=604800)
    return response

@app.get("/signup", response_class=HTMLResponse)
async def signup_page(request: Request):
    return templates.TemplateResponse(request, "signup.html")

@app.post("/signup")
async def signup(request: Request, username: str = Form(...), email: str = Form(...), 
                password: str = Form(...), full_name: str = Form(...), role: str = Form("student")):
    db = get_db()
    try:
        existing = db.execute("SELECT id FROM users WHERE username=? OR email=?", (username, email)).fetchone()
        if existing:
            return templates.TemplateResponse(request, "signup.html", {"error": "Username or email already exists"})
        db.execute(
            "INSERT INTO users (username, email, password_hash, role, full_name) VALUES (?, ?, ?, ?, ?)",
            (username, email, hash_password(password), role, full_name)
        )
        db.commit()
        user = db.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
        token = create_session(user['id'])
        db.close()
        response = RedirectResponse(url="/dashboard", status_code=303)
        response.set_cookie("session_token", token, httponly=True, max_age=604800)
        return response
    except Exception as e:
        db.close()
        return templates.TemplateResponse(request, "signup.html", {"error": str(e)})

@app.get("/logout")
async def logout(request: Request):
    token = request.cookies.get("session_token")
    if token:
        db = get_db()
        db.execute("DELETE FROM sessions WHERE token=?", (token,))
        db.commit()
        db.close()
    response = RedirectResponse(url="/", status_code=303)
    response.delete_cookie("session_token")
    return response

# ─── DASHBOARD ───
@app.get("/dashboard", response_class=HTMLResponse)
async def dashboard(request: Request):
    user = get_current_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=303)
    
    db = get_db()
    stats = {
        'teachers': db.execute("SELECT COUNT(*) as cnt FROM teachers").fetchone()['cnt'],
        'students': db.execute("SELECT COUNT(*) as cnt FROM users WHERE role='student'").fetchone()['cnt'],
        'classes': db.execute("SELECT COUNT(*) as cnt FROM classes").fetchone()['cnt'],
        'subjects': db.execute("SELECT COUNT(*) as cnt FROM subjects").fetchone()['cnt'],
        'timetable_entries': db.execute("SELECT COUNT(*) as cnt FROM timetable").fetchone()['cnt'],
    }
    config = dict(db.execute("SELECT * FROM config WHERE id=1").fetchone())
    db.close()
    
    return templates.TemplateResponse(request, "dashboard.html", {
        "user": user, "stats": stats, "config": config
    })

# ─── TEACHERS CRUD ───
@app.get("/api/teachers")
async def list_teachers(request: Request):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    db = get_db()
    teachers = [dict(r) for r in db.execute("SELECT * FROM teachers ORDER BY name").fetchall()]
    for t in teachers:
        subs = db.execute(
            "SELECT s.id, s.name FROM subjects s JOIN teacher_subjects ts ON s.id=ts.subject_id WHERE ts.teacher_id=?",
            (t['id'],)
        ).fetchall()
        t['subjects'] = [dict(s) for s in subs]
    db.close()
    return teachers

@app.post("/api/teachers")
async def create_teacher(request: Request, name: str = Form(...), employee_id: str = Form(""),
                         specialization: str = Form(""), max_periods_per_day: int = Form(6),
                         phone: str = Form(""), email: str = Form(""), subject_ids: str = Form("[]")):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    db = get_db()
    try:
        if not employee_id or employee_id.strip() == '':
            max_id = db.execute("SELECT COALESCE(MAX(id),0)+1 FROM teachers").fetchone()[0]
            employee_id = f"T{max_id:03d}"
        existing = db.execute("SELECT id FROM teachers WHERE employee_id=?", (employee_id,)).fetchone()
        if existing:
            db.close()
            raise HTTPException(400, detail=f"Employee ID '{employee_id}' already exists")
        cursor = db.execute(
            "INSERT INTO teachers (name, employee_id, specialization, max_periods_per_day, phone, email) VALUES (?,?,?,?,?,?)",
            (name, employee_id, specialization, max_periods_per_day, phone, email)
        )
        teacher_id = cursor.lastrowid
        subjects = json.loads(subject_ids)
        for sid in subjects:
            db.execute("INSERT INTO teacher_subjects (teacher_id, subject_id) VALUES (?,?)", (teacher_id, sid))
        
        for day in range(6):
            for period in range(8):
                db.execute(
                    "INSERT OR IGNORE INTO teacher_availability (teacher_id, day_of_week, period_number, is_available) VALUES (?,?,?,1)",
                    (teacher_id, day, period)
                )
        db.commit()
        return {"success": True, "id": teacher_id}
    except Exception as e:
        db.close()
        raise HTTPException(400, str(e))

@app.put("/api/teachers/{teacher_id}")
async def update_teacher(request: Request, teacher_id: int, name: str = Form(...), 
                         employee_id: str = Form(""), specialization: str = Form(""),
                         max_periods_per_day: int = Form(6), phone: str = Form(""),
                         email: str = Form(""), subject_ids: str = Form("[]")):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    db = get_db()
    db.execute(
        "UPDATE teachers SET name=?, employee_id=?, specialization=?, max_periods_per_day=?, phone=?, email=? WHERE id=?",
        (name, employee_id, specialization, max_periods_per_day, phone, email, teacher_id)
    )
    db.execute("DELETE FROM teacher_subjects WHERE teacher_id=?", (teacher_id,))
    subjects = json.loads(subject_ids)
    for sid in subjects:
        db.execute("INSERT INTO teacher_subjects (teacher_id, subject_id) VALUES (?,?)", (teacher_id, sid))
    db.commit()
    db.close()
    return {"success": True}

@app.delete("/api/teachers/{teacher_id}")
async def delete_teacher(request: Request, teacher_id: int):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    db = get_db()
    db.execute("DELETE FROM teachers WHERE id=?", (teacher_id,))
    db.commit()
    db.close()
    return {"success": True}

# ─── SUBJECTS CRUD ───
@app.get("/api/subjects")
async def list_subjects(request: Request):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    db = get_db()
    subjects = [dict(r) for r in db.execute("SELECT * FROM subjects ORDER BY name").fetchall()]
    db.close()
    return subjects

@app.post("/api/subjects")
async def create_subject(request: Request, name: str = Form(...), code: str = Form(...),
                         color: str = Form("#3b82f6"), periods_per_week: int = Form(5),
                         needs_lab: bool = Form(False)):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    db = get_db()
    try:
        cursor = db.execute(
            "INSERT INTO subjects (name, code, color, periods_per_week, needs_lab) VALUES (?,?,?,?,?)",
            (name, code, color, periods_per_week, 1 if needs_lab else 0)
        )
        db.commit()
        return {"success": True, "id": cursor.lastrowid}
    except Exception as e:
        db.close()
        raise HTTPException(400, str(e))

@app.put("/api/subjects/{subject_id}")
async def update_subject(request: Request, subject_id: int, name: str = Form(...),
                         code: str = Form(...), color: str = Form("#3b82f6"),
                         periods_per_week: int = Form(5), needs_lab: bool = Form(False)):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    db = get_db()
    db.execute(
        "UPDATE subjects SET name=?, code=?, color=?, periods_per_week=?, needs_lab=? WHERE id=?",
        (name, code, color, periods_per_week, 1 if needs_lab else 0, subject_id)
    )
    db.commit()
    db.close()
    return {"success": True}

@app.delete("/api/subjects/{subject_id}")
async def delete_subject(request: Request, subject_id: int):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    db = get_db()
    db.execute("DELETE FROM subjects WHERE id=?", (subject_id,))
    db.commit()
    db.close()
    return {"success": True}

# ─── CLASSES CRUD ───
@app.get("/api/classes")
async def list_classes(request: Request):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    db = get_db()
    classes = [dict(r) for r in db.execute("SELECT * FROM classes ORDER BY grade, section").fetchall()]
    for c in classes:
        subs = db.execute(
            """SELECT s.id, s.name, s.code, cs.periods_per_week 
               FROM subjects s JOIN class_subjects cs ON s.id=cs.subject_id WHERE cs.class_id=?""",
            (c['id'],)
        ).fetchall()
        c['subjects'] = [dict(s) for s in subs]
    db.close()
    return classes

def init_class_time_slots(db, class_id):
    """Generate default time slots for a class based on global config."""
    config = db.execute("SELECT * FROM config WHERE id=1").fetchone()
    if not config:
        return
    periods = config['periods_per_day']
    start = config['start_time']
    duration = config['period_duration']
    lunch = config['lunch_period']
    lunch_dur = config['lunch_duration']
    
    h, m = map(int, start.split(':'))
    base_minutes = h * 60 + m
    gap = 5  # minutes between periods
    
    for p in range(periods):
        s_min = base_minutes + p * (duration + gap)
        if p == lunch:
            e_min = s_min + lunch_dur
        else:
            e_min = s_min + duration
        sh, sm = divmod(s_min, 60)
        eh, em = divmod(e_min, 60)
        db.execute(
            """INSERT OR IGNORE INTO class_time_slots (class_id, period_number, start_time, end_time, is_break)
               VALUES (?, ?, ?, ?, ?)""",
            (class_id, p, f"{sh:02d}:{sm:02d}", f"{eh:02d}:{em:02d}", 1 if p == lunch else 0)
        )

@app.post("/api/classes")
async def create_class(request: Request, name: str = Form(...), grade: str = Form(...),
                       section: str = Form(...), strength: int = Form(40), room_id: int = Form(None),
                       subject_data: str = Form("{}")):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    db = get_db()
    try:
        cursor = db.execute(
            "INSERT INTO classes (name, grade, section, strength, room_id) VALUES (?,?,?,?,?)",
            (name, grade, section, strength, room_id)
        )
        class_id = cursor.lastrowid
        subjects = json.loads(subject_data)
        for sid, periods in subjects.items():
            db.execute("INSERT INTO class_subjects (class_id, subject_id, periods_per_week) VALUES (?,?,?)",
                       (class_id, int(sid), periods))
        init_class_time_slots(db, class_id)
        db.commit()
        return {"success": True, "id": class_id}
    except Exception as e:
        db.close()
        raise HTTPException(400, str(e))

@app.put("/api/classes/{class_id}")
async def update_class(request: Request, class_id: int, name: str = Form(...),
                       grade: str = Form(...), section: str = Form(...), strength: int = Form(40),
                       room_id: int = Form(None), subject_data: str = Form("{}")):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    db = get_db()
    db.execute(
        "UPDATE classes SET name=?, grade=?, section=?, strength=?, room_id=? WHERE id=?",
        (name, grade, section, strength, room_id, class_id)
    )
    db.execute("DELETE FROM class_subjects WHERE class_id=?", (class_id,))
    subjects = json.loads(subject_data)
    for sid, periods in subjects.items():
        db.execute("INSERT INTO class_subjects (class_id, subject_id, periods_per_week) VALUES (?,?,?)",
                   (class_id, int(sid), periods))
    db.commit()
    db.close()
    return {"success": True}

@app.delete("/api/classes/{class_id}")
async def delete_class(request: Request, class_id: int):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    db = get_db()
    db.execute("DELETE FROM classes WHERE id=?", (class_id,))
    db.commit()
    db.close()
    return {"success": True}

# ─── CLASS TIME SLOTS ───
@app.get("/api/classes/{class_id}/time-slots")
async def get_class_time_slots(request: Request, class_id: int):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    db = get_db()
    slots = db.execute(
        "SELECT * FROM class_time_slots WHERE class_id=? ORDER BY period_number",
        (class_id,)
    ).fetchall()
    # If no slots exist yet, generate defaults
    if not slots:
        init_class_time_slots(db, class_id)
        db.commit()
        slots = db.execute(
            "SELECT * FROM class_time_slots WHERE class_id=? ORDER BY period_number",
            (class_id,)
        ).fetchall()
    db.close()
    return [dict(s) for s in slots]

@app.put("/api/classes/{class_id}/time-slots")
async def update_class_time_slots(request: Request, class_id: int):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    body = await request.json()
    db = get_db()
    db.execute("DELETE FROM class_time_slots WHERE class_id=?", (class_id,))
    for slot in body:
        db.execute(
            """INSERT INTO class_time_slots (class_id, period_number, start_time, end_time, is_break)
               VALUES (?, ?, ?, ?, ?)""",
            (class_id, slot['period_number'], slot['start_time'], slot['end_time'], slot.get('is_break', 0))
        )
    db.commit()
    db.close()
    return {"success": True}

@app.post("/api/classes/{class_id}/time-slots/reset")
async def reset_class_time_slots(request: Request, class_id: int):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    db = get_db()
    db.execute("DELETE FROM class_time_slots WHERE class_id=?", (class_id,))
    init_class_time_slots(db, class_id)
    db.commit()
    db.close()
    return {"success": True}

# ─── ROOMS CRUD ───
@app.get("/api/rooms")
async def list_rooms(request: Request):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    db = get_db()
    rooms = [dict(r) for r in db.execute("SELECT * FROM rooms ORDER BY name").fetchall()]
    db.close()
    return rooms

@app.post("/api/rooms")
async def create_room(request: Request, name: str = Form(...), type: str = Form("classroom"),
                      capacity: int = Form(40), equipment: str = Form("")):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    db = get_db()
    cursor = db.execute(
        "INSERT INTO rooms (name, type, capacity, equipment) VALUES (?,?,?,?)",
        (name, type, capacity, equipment)
    )
    db.commit()
    db.close()
    return {"success": True, "id": cursor.lastrowid}

@app.delete("/api/rooms/{room_id}")
async def delete_room(request: Request, room_id: int):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    db = get_db()
    db.execute("DELETE FROM rooms WHERE id=?", (room_id,))
    db.commit()
    db.close()
    return {"success": True}

# ─── CONFIG ───
@app.get("/api/config")
async def get_config(request: Request):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    db = get_db()
    config = dict(db.execute("SELECT * FROM config WHERE id=1").fetchone())
    db.close()
    return config

@app.post("/api/config")
async def update_config(request: Request, school_name: str = Form("My School"),
                        periods_per_day: int = Form(8), period_duration: int = Form(45),
                        start_time: str = Form("08:00"), working_days: str = Form("0,1,2,3,4,5"),
                        lunch_period: int = Form(4), lunch_duration: int = Form(30)):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    db = get_db()
    db.execute(
        "UPDATE config SET school_name=?, periods_per_day=?, period_duration=?, start_time=?, working_days=?, lunch_period=?, lunch_duration=? WHERE id=1",
        (school_name, periods_per_day, period_duration, start_time, working_days, lunch_period, lunch_duration)
    )
    db.commit()
    db.close()
    return {"success": True}

# ─── TIMETABLE GENERATION ───
@app.post("/api/timetable/generate")
async def generate_timetable(request: Request):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    generator = TimetableGenerator()
    success = generator.generate()
    if success:
        return {"success": True, "message": "Timetable generated successfully!"}
    else:
        return JSONResponse(status_code=400, content={"success": False, "errors": generator.errors})

@app.get("/api/timetable")
async def get_timetable(request: Request, class_id: int = None, teacher_id: int = None):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    db = get_db()
    
    if class_id:
        rows = db.execute(
            """SELECT t.*, s.name as subject_name, s.code as subject_code, s.color as subject_color,
                      te.name as teacher_name, r.name as room_name
               FROM timetable t
               JOIN subjects s ON t.subject_id = s.id
               JOIN teachers te ON t.teacher_id = te.id
               LEFT JOIN rooms r ON t.room_id = r.id
               WHERE t.class_id = ?
               ORDER BY t.day_of_week, t.period_number""",
            (class_id,)
        ).fetchall()
    elif teacher_id:
        rows = db.execute(
            """SELECT t.*, s.name as subject_name, s.code as subject_code, s.color as subject_color,
                      c.name as class_name, r.name as room_name
               FROM timetable t
               JOIN subjects s ON t.subject_id = s.id
               JOIN classes c ON t.class_id = c.id
               LEFT JOIN rooms r ON t.room_id = r.id
               WHERE t.teacher_id = ?
               ORDER BY t.day_of_week, t.period_number""",
            (teacher_id,)
        ).fetchall()
    else:
        rows = db.execute(
            """SELECT t.*, s.name as subject_name, s.code as subject_code, s.color as subject_color,
                      te.name as teacher_name, c.name as class_name, r.name as room_name
               FROM timetable t
               JOIN subjects s ON t.subject_id = s.id
               JOIN teachers te ON t.teacher_id = te.id
               JOIN classes c ON t.class_id = c.id
               LEFT JOIN rooms r ON t.room_id = r.id
               ORDER BY t.class_id, t.day_of_week, t.period_number"""
        ).fetchall()
    
    db.close()
    return [dict(r) for r in rows]

@app.post("/api/timetable/swap")
async def swap_periods(request: Request):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    data = await request.json()
    generator = TimetableGenerator()
    generator.load_data()
    success, message = generator.swap_periods(
        data['class_id'], data['day1'], data['period1'], data['day2'], data['period2']
    )
    if success:
        return {"success": True, "message": message}
    else:
        raise HTTPException(400, detail={"success": False, "message": message})

@app.get("/api/timetable/conflicts")
async def check_conflicts(request: Request):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    db = get_db()
    teacher_conflicts = db.execute(
        """SELECT t.day_of_week, t.period_number, te.name, COUNT(*) as cnt
           FROM timetable t JOIN teachers te ON t.teacher_id = te.id
           GROUP BY t.teacher_id, t.day_of_week, t.period_number
           HAVING cnt > 1"""
    ).fetchall()
    db.close()
    return {"conflicts": [dict(r) for r in teacher_conflicts], "count": len(teacher_conflicts)}

# ─── AVAILABILITY ───
@app.get("/api/availability/{teacher_id}")
async def get_availability(request: Request, teacher_id: int):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    db = get_db()
    rows = db.execute(
        "SELECT * FROM teacher_availability WHERE teacher_id=?", (teacher_id,)
    ).fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.post("/api/availability/{teacher_id}")
async def update_availability(request: Request, teacher_id: int):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    data = await request.json()
    db = get_db()
    db.execute("DELETE FROM teacher_availability WHERE teacher_id=?", (teacher_id,))
    for entry in data['availability']:
        db.execute(
            "INSERT INTO teacher_availability (teacher_id, day_of_week, period_number, is_available) VALUES (?,?,?,?)",
            (teacher_id, entry['day'], entry['period'], 1 if entry['available'] else 0)
        )
    db.commit()
    db.close()
    return {"success": True}

# ─── SUBSTITUTION ───
@app.post("/api/substitution")
async def create_substitution(request: Request):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    data = await request.json()
    db = get_db()
    
    gen = TimetableGenerator()
    gen.load_data()
    gen.load_data()
    
    substitute = gen.find_substitute(data['teacher_id'], data['day'], data['period'])
    if substitute:
        db.execute(
            """INSERT INTO substitutions (original_teacher_id, substitute_teacher_id, class_id, subject_id, day_of_week, period_number, date, status)
               VALUES (?,?,?,?,?,?,?,?)""",
            (data['teacher_id'], substitute['id'], data['class_id'], data['subject_id'],
             data['day'], data['period'], data.get('date', ''), 'approved')
        )
        db.commit()
        db.close()
        return {"success": True, "substitute": substitute['name']}
    else:
        db.close()
        return {"success": False, "message": "No available substitute found"}

@app.get("/api/substitutions")
async def list_substitutions(request: Request):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    db = get_db()
    rows = db.execute(
        """SELECT sub.*, t1.name as original_teacher, t2.name as substitute_teacher,
                  c.name as class_name, s.name as subject_name
           FROM substitutions sub
           JOIN teachers t1 ON sub.original_teacher_id = t1.id
           LEFT JOIN teachers t2 ON sub.substitute_teacher_id = t2.id
           JOIN classes c ON sub.class_id = c.id
           JOIN subjects s ON sub.subject_id = s.id
           ORDER BY sub.created_at DESC"""
    ).fetchall()
    db.close()
    return [dict(r) for r in rows]

# ─── EXPORT ───
@app.get("/api/export/csv")
async def export_csv(request: Request, class_id: int = None):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    db = get_db()
    
    if class_id:
        rows = db.execute(
            """SELECT t.day_of_week, t.period_number, s.name as subject, te.name as teacher, r.name as room
               FROM timetable t JOIN subjects s ON t.subject_id=s.id JOIN teachers te ON t.teacher_id=te.id
               LEFT JOIN rooms r ON t.room_id=r.id WHERE t.class_id=? ORDER BY t.day_of_week, t.period_number""",
            (class_id,)
        ).fetchall()
    else:
        rows = db.execute(
            """SELECT c.name as class_name, t.day_of_week, t.period_number, s.name as subject, 
                      te.name as teacher, r.name as room
               FROM timetable t JOIN subjects s ON t.subject_id=s.id JOIN teachers te ON t.teacher_id=te.id
               JOIN classes c ON t.class_id=c.id LEFT JOIN rooms r ON t.room_id=r.id
               ORDER BY c.name, t.day_of_week, t.period_number"""
        ).fetchall()
    
    db.close()
    
    output = io.StringIO()
    writer = csv.writer(output)
    if class_id:
        writer.writerow(['Day', 'Period', 'Subject', 'Teacher', 'Room'])
    else:
        writer.writerow(['Class', 'Day', 'Period', 'Subject', 'Teacher', 'Room'])
    
    for row in rows:
        if class_id:
            writer.writerow([DAYS[row['day_of_week']], row['period_number']+1, 
                           row['subject'], row['teacher'], row['room']])
        else:
            writer.writerow([row['class_name'], DAYS[row['day_of_week']], row['period_number']+1,
                           row['subject'], row['teacher'], row['room']])
    
    output.seek(0)
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode()),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=timetable.csv"}
    )

@app.get("/api/export/excel")
async def export_excel(request: Request, class_id: int = None):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    
    db = get_db()
    if class_id:
        rows = db.execute(
            """SELECT t.day_of_week, t.period_number, s.name as subject, s.color, 
                      te.name as teacher, r.name as room
               FROM timetable t JOIN subjects s ON t.subject_id=s.id JOIN teachers te ON t.teacher_id=te.id
               LEFT JOIN rooms r ON t.room_id=r.id WHERE t.class_id=? ORDER BY t.day_of_week, t.period_number""",
            (class_id,)
        ).fetchall()
    else:
        rows = db.execute(
            """SELECT c.name as class_name, t.day_of_week, t.period_number, s.name as subject, s.color,
                      te.name as teacher, r.name as room
               FROM timetable t JOIN subjects s ON t.subject_id=s.id JOIN teachers te ON t.teacher_id=te.id
               JOIN classes c ON t.class_id=c.id LEFT JOIN rooms r ON t.room_id=r.id
               ORDER BY c.name, t.day_of_week, t.period_number"""
        ).fetchall()
    db.close()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Timetable"
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="10b981", end_color="10b981", fill_type="solid")
    
    if class_id:
        ws.append(['Day', 'Period', 'Subject', 'Teacher', 'Room'])
        for col in range(1, 6):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
        for row in rows:
            ws.append([DAYS[row['day_of_week']], row['period_number']+1,
                      row['subject'], row['teacher'], row['room']])
    else:
        ws.append(['Class', 'Day', 'Period', 'Subject', 'Teacher', 'Room'])
        for col in range(1, 7):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
        for row in rows:
            ws.append([row['class_name'], DAYS[row['day_of_week']], row['period_number']+1,
                      row['subject'], row['teacher'], row['room']])
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=timetable.xlsx"}
    )

@app.get("/api/export/pdf")
async def export_pdf(request: Request, class_id: int = None):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)

    from fpdf import FPDF

    db = get_db()
    if class_id:
        rows = db.execute(
            """SELECT t.day_of_week, t.period_number, s.name as subject, s.color as subject_color,
                      te.name as teacher, r.name as room, c.name as class_name
               FROM timetable t JOIN subjects s ON t.subject_id=s.id JOIN teachers te ON t.teacher_id=te.id
               LEFT JOIN rooms r ON t.room_id=r.id JOIN classes c ON t.class_id=c.id
               WHERE t.class_id=? ORDER BY t.day_of_week, t.period_number""",
            (class_id,)
        ).fetchall()
    else:
        rows = db.execute(
            """SELECT c.name as class_name, t.day_of_week, t.period_number, s.name as subject, s.color as subject_color,
                      te.name as teacher, r.name as room
               FROM timetable t JOIN subjects s ON t.subject_id=s.id JOIN teachers te ON t.teacher_id=te.id
               JOIN classes c ON t.class_id=c.id LEFT JOIN rooms r ON t.room_id=r.id
               ORDER BY c.name, t.day_of_week, t.period_number"""
        ).fetchall()

    config = db.execute("SELECT * FROM config WHERE id=1").fetchone()
    db.close()

    school_name = config['school_name'] if config else 'Timetable'
    periods_per_day = config['periods_per_day'] if config else 8

    # Group by class
    classes_data = {}
    for r in rows:
        cls_name = r['class_name']
        if cls_name not in classes_data:
            classes_data[cls_name] = {}
        day = r['day_of_week']
        if day not in classes_data[cls_name]:
            classes_data[cls_name][day] = {}
        color = r['subject_color'] or '#4F46E5'
        r_hex = int(color[1:3], 16) if len(color) >= 7 else 79
        g_hex = int(color[3:5], 16) if len(color) >= 7 else 70
        b_hex = int(color[5:7], 16) if len(color) >= 7 else 229
        classes_data[cls_name][day][r['period_number']] = {
            'text': f"{r['subject']} - {r['teacher']}" + (f" ({r['room']})" if r['room'] else ""),
            'color': (r_hex, g_hex, b_hex)
        }

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)

    for cls_name, days in classes_data.items():
        pdf.add_page()
        # Header
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, school_name, new_x="LMARGIN", new_y="NEXT", align="C")
        pdf.set_font("Helvetica", "", 12)
        pdf.cell(0, 8, f"Class: {cls_name}", new_x="LMARGIN", new_y="NEXT", align="C")
        pdf.ln(5)

        # Table header
        pdf.set_font("Helvetica", "B", 9)
        col_w = 180 // (periods_per_day + 1) if (periods_per_day + 1) > 0 else 25
        day_w = 180 - col_w * periods_per_day
        if day_w < 20: day_w = 20; col_w = 160 // periods_per_day

        pdf.set_fill_color(79, 70, 229)
        pdf.set_text_color(255, 255, 255)
        pdf.cell(day_w, 7, "Day / Period", border=1, fill=True)
        for p in range(periods_per_day):
            pdf.cell(col_w, 7, f"P{p+1}", border=1, fill=True, align="C")
        pdf.ln()

        # Table rows
        pdf.set_font("Helvetica", "", 7)
        day_names = ['Monday','Tuesday','Wednesday','Thursday','Friday','Saturday']
        for d_idx, d_name in enumerate(day_names):
            if d_idx not in days:
                continue
            pdf.set_text_color(0, 0, 0)
            pdf.set_fill_color(249, 250, 251)
            pdf.cell(day_w, 7, d_name[:4], border=1, fill=True)
            for p in range(periods_per_day):
                cell = days[d_idx].get(p)
                if cell:
                    r, g, b = cell['color']
                    pdf.set_text_color(r, g, b)
                    lr = min(255, r + 160)
                    lg = min(255, g + 160)
                    lb = min(255, b + 160)
                    pdf.set_fill_color(lr, lg, lb)
                    pdf.set_font("Helvetica", "B", 7)
                    pdf.cell(col_w, 7, cell['text'][:20], border=1, align="C", fill=True)
                    pdf.set_font("Helvetica", "", 7)
                    pdf.set_text_color(0, 0, 0)
                else:
                    pdf.cell(col_w, 7, "", border=1, align="C")
            pdf.ln()

    output = io.BytesIO()
    pdf.output(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=timetable.pdf"}
    )


# ─── NOTIFICATIONS ───
@app.get("/api/notifications")
async def list_notifications(request: Request):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    db = get_db()
    rows = db.execute(
        "SELECT * FROM notifications WHERE user_id=? ORDER BY created_at DESC LIMIT 50",
        (user['id'],)
    ).fetchall()
    db.close()
    return [dict(r) for r in rows]

@app.post("/api/notifications/read")
async def mark_notifications_read(request: Request):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    db = get_db()
    db.execute("UPDATE notifications SET is_read=1 WHERE user_id=?", (user['id'],))
    db.commit()
    db.close()
    return {"success": True}

# ─── BULK IMPORT ───
@app.post("/api/import/teachers")
async def import_teachers(request: Request, file: UploadFile = File(...)):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    content = await file.read().decode('utf-8')
    reader = csv.DictReader(io.StringIO(content))
    db = get_db()
    count = 0
    for row in reader:
        try:
            db.execute(
                "INSERT INTO teachers (name, employee_id, specialization, max_periods_per_day) VALUES (?,?,?,?)",
                (row.get('name', ''), row.get('employee_id', ''), row.get('specialization', ''), int(row.get('max_periods', 6)))
            )
            count += 1
        except:
            continue
    db.commit()
    db.close()
    return {"success": True, "imported": count}

@app.post("/api/import/subjects")
async def import_subjects(request: Request, file: UploadFile = File(...)):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    content = await file.read().decode('utf-8')
    reader = csv.DictReader(io.StringIO(content))
    db = get_db()
    count = 0
    for row in reader:
        try:
            db.execute(
                "INSERT INTO subjects (name, code, periods_per_week, needs_lab) VALUES (?,?,?,?)",
                (row.get('name', ''), row.get('code', ''), int(row.get('periods', 5)), 1 if row.get('needs_lab') == 'true' else 0)
            )
            count += 1
        except:
            continue
    db.commit()
    db.close()
    return {"success": True, "imported": count}

@app.post("/api/import/classes")
async def import_classes(request: Request, file: UploadFile = File(...)):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        raise HTTPException(403)
    content = await file.read().decode('utf-8')
    reader = csv.DictReader(io.StringIO(content))
    db = get_db()
    count = 0
    for row in reader:
        try:
            db.execute(
                "INSERT INTO classes (name, grade, section, strength) VALUES (?,?,?,?)",
                (row.get('name', ''), row.get('grade', ''), row.get('section', ''), int(row.get('strength', 40)))
            )
            count += 1
        except:
            continue
    db.commit()
    db.close()
    return {"success": True, "imported": count}

# ─── DATA ENDPOINT ───
@app.get("/api/data")
async def get_all_data(request: Request):
    user = get_current_user(request)
    if not user:
        raise HTTPException(401)
    db = get_db()
    data = {
        'config': dict(db.execute("SELECT * FROM config WHERE id=1").fetchone()),
        'subjects': [dict(r) for r in db.execute("SELECT * FROM subjects").fetchall()],
        'teachers': [dict(r) for r in db.execute("SELECT * FROM teachers").fetchall()],
        'classes': [dict(r) for r in db.execute("SELECT * FROM classes").fetchall()],
        'rooms': [dict(r) for r in db.execute("SELECT * FROM rooms").fetchall()],
    }
    db.close()
    return data

# ─── PAGES ───
@app.get("/manage/teachers", response_class=HTMLResponse)
async def manage_teachers_page(request: Request):
    user = get_current_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=303)
    return templates.TemplateResponse(request, "manage_teachers.html", {"user": user})

@app.get("/manage/subjects", response_class=HTMLResponse)
async def manage_subjects_page(request: Request):
    user = get_current_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=303)
    return templates.TemplateResponse(request, "manage_subjects.html", {"user": user})

@app.get("/manage/classes", response_class=HTMLResponse)
async def manage_classes_page(request: Request):
    user = get_current_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=303)
    return templates.TemplateResponse(request, "manage_classes.html", {"user": user})

@app.get("/manage/rooms", response_class=HTMLResponse)
async def manage_rooms_page(request: Request):
    user = get_current_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=303)
    return templates.TemplateResponse(request, "manage_rooms.html", {"user": user})

@app.get("/timetable", response_class=HTMLResponse)
async def timetable_page(request: Request, class_id: int = None, teacher_id: int = None):
    user = get_current_user(request)
    if not user:
        return RedirectResponse(url="/login", status_code=303)
    db = get_db()
    classes = [dict(r) for r in db.execute("SELECT * FROM classes ORDER BY grade, section").fetchall()]
    teachers = [dict(r) for r in db.execute("SELECT * FROM teachers ORDER BY name").fetchall()]
    db.close()
    return templates.TemplateResponse(request, "timetable.html", {
        "user": user, "classes": classes, "teachers": teachers,
        "selected_class": class_id, "selected_teacher": teacher_id
    })

@app.get("/settings", response_class=HTMLResponse)
async def settings_page(request: Request):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        return RedirectResponse(url="/login", status_code=303)
    db = get_db()
    config = dict(db.execute("SELECT * FROM config WHERE id=1").fetchone())
    db.close()
    return templates.TemplateResponse(request, "settings.html", {"user": user, "config": config})

@app.post("/settings")
async def save_settings(request: Request):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        return RedirectResponse(url="/login", status_code=303)
    form = await request.form()
    school_name = form.get('school_name', 'My School')
    periods_per_day = int(form.get('periods_per_day', 8))
    period_duration = int(form.get('period_duration', 45))
    start_time = form.get('start_time', '08:00')
    lunch_period = int(form.get('lunch_period', 4))
    
    working_days_values = form.getlist('working_days_values')
    if working_days_values:
        working_days = ','.join(working_days_values)
    else:
        working_days = form.get('working_days', '0,1,2,3,4,5')
    
    db = get_db()
    db.execute(
        "UPDATE config SET school_name=?, periods_per_day=?, period_duration=?, start_time=?, working_days=?, lunch_period=? WHERE id=1",
        (school_name, periods_per_day, period_duration, start_time, working_days, lunch_period)
    )
    db.commit()
    db.close()
    return RedirectResponse(url="/settings", status_code=303)

@app.get("/substitutions", response_class=HTMLResponse)
async def substitutions_page(request: Request):
    user = get_current_user(request)
    if not user or user['role'] != 'admin':
        return RedirectResponse(url="/login", status_code=303)
    return templates.TemplateResponse(request, "substitutions.html", {"user": user})

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
